"""
Targeted verification: just check the Q2/Q3 fill bullets (2 bullets each, right after Script:)
and that no placeholders exist.
"""
from docx import Document

all_guides = [
    ("Lance PM (full)", "guides/lance-pm/Lance_PM_Interview_Prep_Guide.docx"),
    ("MediaAlpha TPM (full)", "guides/mediaalpha-tpm/MediaAlpha_TPM_Interview_Prep_Guide.docx"),
    ("New Relic PM (full)", "guides/new-relic-product-manager-log-management/New_Relic_PM_Interview_Prep_Guide.docx"),
    ("Scale AI PM-TE (full)", "guides/scale-ai-pm-te/ScaleAI_Interview_Prep_Guide.docx"),
    ("Scale AI TPM-CV (full)", "guides/scale-ai-tpm-cv/Scale_AI_Interview_Prep_Guide.docx"),
    ("Podium PM (full)", "guides/podium-pm/Podium_Interview_Prep_Guide.docx"),
    ("Mintlify SE (full)", "guides/mintlify-solutions-engineer-post-sales/Mintlify_SE_PostSales_Interview_Prep_Guide.docx"),
    ("Datadog PTSE (full)", "guides/datadog-partner-technology-solutions-engineer/Datadog_PTSE_Interview_Prep_Guide.docx"),
    ("Lance PM (short)", "guides/lance-pm/Lance_PM_Shorter_Interview_Prep_Guide.docx"),
    ("MediaAlpha TPM (short)", "guides/mediaalpha-tpm/MediaAlpha_TPM_Shorter_Interview_Prep_Guide.docx"),
    ("New Relic PM (short)", "guides/new-relic-product-manager-log-management/New_Relic_PM_Shorter_Interview_Prep_Guide.docx"),
    ("Scale AI PM-TE (short)", "guides/scale-ai-pm-te/ScaleAI_Interview_Prep_Guide_Short.docx"),
    ("Scale AI TPM-CV (short)", "guides/scale-ai-tpm-cv/Scale_AI_Interview_Prep_Guide_Short.docx"),
    ("Podium PM (short)", "guides/podium-pm/Podium_Interview_Prep_Guide_Short.docx"),
]

def check_guide(path):
    doc = Document(path)
    paras = doc.paragraphs
    issues = []

    # Find Q2 and Q3 section markers
    q2_idx = None
    q3_idx = None
    for i, p in enumerate(paras):
        txt = p.text.strip()
        if p.style.name.startswith('Heading') and 'Q2' in txt and ('Why' in txt or 'here' in txt.lower()):
            q2_idx = i
        if p.style.name.startswith('Heading') and 'Q3' in txt and ('Do you know' in txt or 'What do we do' in txt or 'what this role' in txt.lower()):
            q3_idx = i

    if q2_idx is None or q3_idx is None:
        return [f"Could not find Q2/Q3 headings (q2={q2_idx}, q3={q3_idx})"]

    # For each Q, find the Script: line and the next 2 content paragraphs
    def get_fill_bullets(start_idx):
        bullets = []
        found_script = False
        for p in paras[start_idx+1:]:
            if p.style.name.startswith('Heading'):
                break
            txt = p.text.strip()
            if 'Script' in txt and not found_script:
                found_script = True
                continue
            if found_script and txt:
                bullets.append(p)
                if len(bullets) == 2:
                    break
        return bullets

    q2_bullets = get_fill_bullets(q2_idx)
    q3_bullets = get_fill_bullets(q3_idx)

    # Check each bullet: no placeholder, has bold first run with colon
    for label, bullets in [("Q2", q2_bullets), ("Q3", q3_bullets)]:
        if len(bullets) < 2:
            issues.append(f"{label}: only {len(bullets)} bullet(s) found")
            continue
        for b in bullets:
            txt = b.text.strip()
            if 'Fill in here' in txt:
                issues.append(f"{label}: PLACEHOLDER not filled: '{txt[:60]}'")
                continue
            if not b.runs:
                issues.append(f"{label}: No runs in bullet: '{txt[:60]}'")
                continue
            r0 = b.runs[0]
            if not r0.bold:
                issues.append(f"{label}: First run NOT bold: '{r0.text[:40]}'")
            if ':' not in r0.text:
                issues.append(f"{label}: First run has no colon: '{r0.text[:40]}'")

    return issues

all_ok = True
for name, path in all_guides:
    try:
        issues = check_guide(path)
        if issues:
            print(f"FAIL {name}:")
            for iss in issues:
                print(f"  - {iss}")
            all_ok = False
        else:
            print(f"OK   {name}")
    except Exception as err:
        print(f"ERR  {name}: {err}")
        all_ok = False

print(f"\n{'ALL CLEAN' if all_ok else 'ISSUES FOUND'}")
